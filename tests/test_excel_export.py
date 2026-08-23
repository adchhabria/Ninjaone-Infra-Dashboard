"""Tests for Multi-Sheet Excel Report Generator."""

import io
import openpyxl
import pytest

from scripts.generate_sample_data import get_mock_dashboard_data
from src.metrics.excel_export import generate_excel_workbook


class TestExcelExport:
    def test_generate_excel_workbook(self):
        data = get_mock_dashboard_data()
        excel_bytes = generate_excel_workbook(data)

        assert isinstance(excel_bytes, bytes)
        assert len(excel_bytes) > 1000

        # Validate with openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        sheet_names = wb.sheetnames

        assert "Executive Summary" in sheet_names
        assert "Organization Compliance" in sheet_names
        assert "EOL Device Ledger" in sheet_names
        assert "Patch SLA Aging" in sheet_names
        assert "Needs Reboot" in sheet_names
        assert "Patch Failures" in sheet_names

        # Verify Executive Summary has rows
        ws = wb["Executive Summary"]
        assert ws.max_row > 5
        assert ws.cell(row=1, column=1).value == "Metric"
