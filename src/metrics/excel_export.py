"""
Enterprise Multi-Sheet Excel (.xlsx) Report Generator.

Inspired by ninjaone-patch-toolkit:
Generates a structured workbook with 6 formatted sheets:
1. Executive Summary
2. Organization Compliance
3. EOL Device Ledger
4. Patch Inventory & SLA Aging
5. Needs Reboot Watchlist
6. Patch Failures
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.metrics.aggregator import DashboardData


def generate_excel_workbook(data: DashboardData) -> bytes:
    """
    Builds a professional multi-sheet .xlsx workbook from DashboardData.
    Returns the binary content as bytes.
    """
    output = io.BytesIO()

    # 1. Prepare DataFrames
    # Sheet 1: Executive Summary
    summary_data = [
        {"Metric": "Report Title", "Value": "NinjaOne IT Infrastructure & Patch Compliance Audit"},
        {"Metric": "Generated Timestamp (UTC)", "Value": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")},
        {"Metric": "Active Scope / Scope Filter", "Value": data.active_filter_label},
        {"Metric": "Total Managed Fleet", "Value": data.total_devices},
        {"Metric": "Reachable Devices (Online)", "Value": f"{data.online_devices} ({data.online_pct:.1f}%)"},
        {"Metric": "Overall Compliance Score", "Value": f"{data.overall_compliance_score:.1f}% ({data.compliance_rag})"},
        {"Metric": "Total Server Infrastructure", "Value": data.total_servers},
        {"Metric": "EOL Devices at Risk", "Value": data.eol_risk_count},
        {"Metric": "Patch Fleet Coverage", "Value": f"{data.patches.get('patch_coverage_pct', 0):.1f}%"},
        {"Metric": "Pending Reboot Count", "Value": data.sla.get("reboot_required_count", 0)},
        {"Metric": "Active Patch Failures", "Value": data.sla.get("failed_patches_count", 0)},
        {"Metric": "SLA Breach (>90 Days)", "Value": data.sla.get("sla_breach_count", 0)},
    ]
    df_summary = pd.DataFrame(summary_data)

    # Sheet 2: Organization Compliance
    df_orgs = pd.DataFrame(data.org_table)
    if not df_orgs.empty:
        df_orgs = df_orgs.rename(columns={
            "org_name": "Organization",
            "region": "Region",
            "country": "Country",
            "device_count": "Devices",
            "online_pct": "Online %",
            "patch_pct": "Patch %",
            "os_pct": "OS Compliance %",
            "eol_count": "EOL Devices",
            "compliance_score": "Score",
            "rag": "Status",
        })
        # Drop internal id column if present
        df_orgs = df_orgs[[c for c in df_orgs.columns if c != "org_id"]]

    # Sheet 3: EOL Ledger
    eol_list = data.os.get("eol_table_data", [])
    df_eol = pd.DataFrame(eol_list)
    if not df_eol.empty:
        df_eol = df_eol.rename(columns={
            "name": "Device Name",
            "org_name": "Organization",
            "region": "Region",
            "os": "Operating System",
            "eol_date": "EOL Date",
            "days_overdue": "Days Overdue",
            "status": "Support Status",
            "risk_level": "Risk Level",
        })
        df_eol = df_eol[[c for c in df_eol.columns if c != "id"]]

    # Sheet 4: Patch Inventory & SLA Aging
    patches_list = data.sla.get("patch_detail_rows", [])
    df_patches = pd.DataFrame(patches_list)
    if not df_patches.empty:
        df_patches = df_patches.rename(columns={
            "device_name": "Device Name",
            "org_name": "Organization",
            "region": "Region",
            "patch_name": "Patch / Update Name",
            "patch_type": "Patch Type",
            "severity": "Severity",
            "age_days": "Age (Days)",
            "sla_bucket": "SLA Aging Bucket",
            "status": "Approval Status",
        })
        df_patches = df_patches[[c for c in df_patches.columns if c != "device_id"]]

    # Sheet 5: Needs Reboot
    reboot_list = data.sla.get("reboot_devices", [])
    df_reboot = pd.DataFrame(reboot_list)
    if not df_reboot.empty:
        df_reboot = df_reboot.rename(columns={
            "name": "Device Name",
            "org_name": "Organization",
            "region": "Region",
            "os": "Operating System",
            "uptime_days": "Uptime (Days)",
            "status": "State",
            "pending_count": "Pending Patches",
        })
        df_reboot = df_reboot[[c for c in df_reboot.columns if c != "device_id"]]

    # Sheet 6: Patch Failures
    failed_list = data.sla.get("failed_patches", [])
    df_failed = pd.DataFrame(failed_list)
    if not df_failed.empty:
        df_failed = df_failed.rename(columns={
            "device_name": "Device Name",
            "org_name": "Organization",
            "region": "Region",
            "patch_name": "Failed Patch / Update",
            "error_code": "Error Code",
            "failed_date": "Failed Date",
            "attempts": "Retry Attempts",
        })
        df_failed = df_failed[[c for c in df_failed.columns if c != "device_id"]]

    # 2. Write to Excel via openpyxl
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_summary.to_excel(writer, sheet_name="Executive Summary", index=False)
        if not df_orgs.empty:
            df_orgs.to_excel(writer, sheet_name="Organization Compliance", index=False)
        if not df_eol.empty:
            df_eol.to_excel(writer, sheet_name="EOL Device Ledger", index=False)
        if not df_patches.empty:
            df_patches.to_excel(writer, sheet_name="Patch SLA Aging", index=False)
        if not df_reboot.empty:
            df_reboot.to_excel(writer, sheet_name="Needs Reboot", index=False)
        if not df_failed.empty:
            df_failed.to_excel(writer, sheet_name="Patch Failures", index=False)

        # 3. Format and Style All Sheets
        wb = writer.book
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="E0E0E0"),
            right=Side(style="thin", color="E0E0E0"),
            top=Side(style="thin", color="E0E0E0"),
            bottom=Side(style="thin", color="E0E0E0"),
        )

        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions

            for col_idx, col in enumerate(ws.columns, 1):
                max_len = 0
                for cell in col:
                    if cell.row == 1:
                        cell.fill = header_fill
                        cell.font = header_font
                        cell.alignment = Alignment(horizontal="center", vertical="center")
                    else:
                        cell.border = thin_border
                        cell.font = Font(name="Calibri", size=10)

                    val_str = str(cell.value or "")
                    max_len = max(max_len, len(val_str))

                col_letter = get_column_letter(col_idx)
                ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    return output.getvalue()
